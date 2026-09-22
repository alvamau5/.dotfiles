#!/usr/bin/env python3
"""
Genera un reporte semanal de obra en Word a partir de reportes diarios (.docx).

Uso:
  python generar_reporte_semanal.py --input ./reportes_diarios --output reporte_semanal_obra.docx
  python generar_reporte_semanal.py --input ./reportes_diarios --baseline linea_base.csv --output reporte.docx

La línea base CSV puede incluir:
  fecha,avance_programado_pct,costo_programado
  2026-05-11,40.0,950000
"""

from __future__ import annotations

import argparse
import csv
import re
import tempfile
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

SPANISH_MONTHS = {
    "enero": 1,
    "febrero": 2,
    "marzo": 3,
    "abril": 4,
    "mayo": 5,
    "junio": 6,
    "julio": 7,
    "agosto": 8,
    "septiembre": 9,
    "setiembre": 9,
    "octubre": 10,
    "noviembre": 11,
    "diciembre": 12,
}

SECTION_RE = re.compile(r"§\s*([+0-9]+)\s+([^\n]+)", re.IGNORECASE)


@dataclass
class DailyReport:
    path: Path
    obra: str = "No identificado"
    no_reporte: str = "No identificado"
    fecha_texto: str = "No identificada"
    fecha: Optional[datetime] = None
    dia_obra: Optional[int] = None
    residente: str = "No identificado"
    director: str = "No identificado"
    resumen: str = ""
    horas_efectivas: Optional[float] = None
    personal_total: Optional[float] = None
    maquinaria: List[Dict[str, str]] = field(default_factory=list)
    actividades: List[Dict[str, str]] = field(default_factory=list)
    materiales: List[Dict[str, str]] = field(default_factory=list)
    incidencias: List[Dict[str, str]] = field(default_factory=list)
    pendientes: List[Dict[str, str]] = field(default_factory=list)
    presupuesto: List[Dict[str, str]] = field(default_factory=list)
    raw_text: str = ""


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip())


def to_float(value: Any) -> Optional[float]:
    if value is None:
        return None
    s = str(value).strip()
    if not s or s in {"—", "-"}:
        return None
    s = s.replace("$", "").replace(",", "").replace("%", "")
    s = re.sub(r"[^0-9.\-]", "", s)
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def parse_spanish_date(text: str) -> Optional[datetime]:
    text = text.lower()
    # Ej: Lunes, 11 de mayo de 2026
    m = re.search(r"(\d{1,2})\s+de\s+([a-záéíóúñ]+)\s+de\s+(\d{4})", text)
    if m:
        day = int(m.group(1))
        month = SPANISH_MONTHS.get(m.group(2).replace("á", "a")) or SPANISH_MONTHS.get(m.group(2))
        year = int(m.group(3))
        if month:
            return datetime(year, month, day)
    # ISO fallback
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", text)
    if m:
        return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


def extract_section(text: str, title_keyword: str) -> str:
    matches = list(SECTION_RE.finditer(text))
    for i, match in enumerate(matches):
        title = match.group(2).lower()
        if title_keyword.lower() in title:
            start = match.end()
            end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
            return text[start:end].strip()
    return ""


def docx_tables(doc: Document) -> List[List[List[str]]]:
    tables = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([normalize_text(cell.text) for cell in row.cells])
        if rows:
            tables.append(rows)
    return tables


def rows_to_dicts(rows: List[List[str]]) -> List[Dict[str, str]]:
    if not rows:
        return []
    header_idx = 0
    # Find first row with at least two non-empty cells as header.
    for i, row in enumerate(rows[:3]):
        if sum(1 for c in row if c.strip()) >= 2:
            header_idx = i
            break
    headers = [normalize_text(h) or f"col_{i+1}" for i, h in enumerate(rows[header_idx])]
    out = []
    for row in rows[header_idx + 1 :]:
        if not any(c.strip() for c in row):
            continue
        item = {headers[i] if i < len(headers) else f"col_{i+1}": normalize_text(c) for i, c in enumerate(row)}
        out.append(item)
    return out


def find_table(tables: List[List[List[str]]], required_keywords: Iterable[str]) -> List[Dict[str, str]]:
    keys = [k.lower() for k in required_keywords]
    for rows in tables:
        flat_header = " ".join(" ".join(row) for row in rows[:2]).lower()
        if all(k in flat_header for k in keys):
            return rows_to_dicts(rows)
    return []


def parse_report(path: Path) -> DailyReport:
    doc = Document(str(path))
    paragraphs = [normalize_text(p.text) for p in doc.paragraphs if normalize_text(p.text)]
    text = "\n".join(paragraphs)
    tables = docx_tables(doc)
    report = DailyReport(path=path, raw_text=text)

    # Datos generales desde texto completo y tablas simples.
    full = text + "\n" + "\n".join(" | ".join(cell for row in table for cell in row) for table in tables)

    patterns = {
        "obra": r"OBRA:\s*([^\n|]+)",
        "no_reporte": r"No\.\s*Reporte:\s*([^\n|]+)",
        "fecha_texto": r"FECHA:\s*([^\n|]+)",
        "residente": r"RESIDENTE:\s*([^\n|]+)",
        "director": r"DIRECTOR:\s*([^\n|]+)",
    }
    for attr, pat in patterns.items():
        m = re.search(pat, full, re.IGNORECASE)
        if m:
            setattr(report, attr, normalize_text(m.group(1)))

    # Fallbacks comunes cuando el texto viene en tablas de dos columnas.
    for table in tables:
        for row in table:
            joined = " | ".join(row)
            if len(row) >= 2:
                label = row[0].strip().lower().rstrip(":")
                value = row[1].strip()
                if label == "obra" and value:
                    report.obra = value
                elif label == "fecha" and value:
                    report.fecha_texto = value
                elif label == "residente" and value:
                    report.residente = value
                elif label == "director" and value:
                    report.director = value
            m = re.search(r"No\.\s*Reporte:\s*([^|]+)", joined, re.IGNORECASE)
            if m:
                report.no_reporte = normalize_text(m.group(1))

    report.fecha = parse_spanish_date(report.fecha_texto or full)
    m = re.search(r"D[ií]a\s+(\d+)\s+de\s+obra", full, re.IGNORECASE)
    if m:
        report.dia_obra = int(m.group(1))

    # Resumen ejecutivo.
    resumen_match = re.search(
        r"RESUMEN EJECUTIVO[^\n]*\n+(.*?)(?=\n+§\s*2|\n+CONDICIONES DEL D[IÍ]A|\Z)",
        text,
        re.IGNORECASE | re.DOTALL,
    )
    if resumen_match:
        report.resumen = normalize_text(resumen_match.group(1))

    # Horas efectivas y personal total.
    m = re.search(r"Horas efectivas:\s*([0-9.,]+)", full, re.IGNORECASE)
    if m:
        report.horas_efectivas = to_float(m.group(1))

    report.maquinaria = find_table(tables, ["Equipo", "Hrs. Op"])
    report.actividades = find_table(tables, ["Clave", "Descripción", "Estado"])
    report.materiales = find_table(tables, ["Material", "Proveedor", "Remisión"])
    report.incidencias = find_table(tables, ["Tipo", "Urgencia", "Descripción"])
    report.pendientes = find_table(tables, ["Tarea", "Responsable", "Fecha"])
    report.presupuesto = find_table(tables, ["Partida", "Presupuestado", "Ejecutado"])

    personal = find_table(tables, ["Cargo", "Cant", "Jornales"])
    if personal:
        for row in personal:
            row_text = " ".join(row.values()).lower()
            if "total personal" in row_text:
                nums = [to_float(v) for v in row.values()]
                nums = [n for n in nums if n is not None]
                if nums:
                    report.personal_total = max(nums)
        if report.personal_total is None:
            total = 0.0
            for row in personal:
                vals = list(row.values())
                if vals and "total" not in vals[0].lower():
                    # Buscar columna Cant.
                    for k, v in row.items():
                        if "cant" in k.lower():
                            n = to_float(v)
                            if n:
                                total += n
            report.personal_total = total if total > 0 else None

    return report


def read_baseline(path: Optional[Path]) -> pd.DataFrame:
    if not path:
        return pd.DataFrame()
    df = pd.read_csv(path)
    if "fecha" in df.columns:
        df["fecha"] = pd.to_datetime(df["fecha"])
    return df


def get_total_budget_row(report: DailyReport) -> Optional[Dict[str, str]]:
    for row in report.presupuesto:
        joined = " ".join(row.values()).lower()
        if "total obra" in joined or joined.strip().startswith("total"):
            return row
    return None


def find_percent_in_row(row: Dict[str, str]) -> Optional[float]:
    for key, value in row.items():
        if "%" in key or "avance" in key.lower():
            val = to_float(value)
            if val is not None:
                return val
    for value in row.values():
        if "%" in str(value):
            val = to_float(value)
            if val is not None:
                return val
    return None


def find_money(row: Dict[str, str], keyword: str) -> Optional[float]:
    for key, value in row.items():
        if keyword.lower() in key.lower():
            return to_float(value)
    return None


def actual_progress_series(reports: List[DailyReport]) -> pd.DataFrame:
    data = []
    for r in reports:
        total = get_total_budget_row(r)
        if total and r.fecha:
            data.append(
                {
                    "fecha": r.fecha,
                    "avance_real_pct": find_percent_in_row(total),
                    "costo_real": find_money(total, "Ejecutado"),
                    "presupuesto_total": find_money(total, "Presupuestado"),
                }
            )
    df = pd.DataFrame(data)
    if not df.empty:
        df = df.sort_values("fecha")
    return df


def make_s_curve(actual: pd.DataFrame, baseline: pd.DataFrame, out_png: Path) -> Tuple[bool, str]:
    if actual.empty:
        return False, "No se encontraron datos suficientes para graficar avance real acumulado."

    plt.figure(figsize=(8, 4.8))

    note = ""
    if "avance_real_pct" in actual and actual["avance_real_pct"].notna().any():
        plt.plot(actual["fecha"], actual["avance_real_pct"], marker="o", label="Avance real acumulado (%)")
    elif "costo_real" in actual and actual["costo_real"].notna().any():
        plt.plot(actual["fecha"], actual["costo_real"], marker="o", label="Costo real acumulado")
        note = "Se graficó costo acumulado porque no se encontró porcentaje de avance real."
    else:
        return False, "No se encontró porcentaje de avance ni costo real acumulado para curva S."

    if not baseline.empty:
        if "avance_programado_pct" in baseline.columns:
            plt.plot(baseline["fecha"], baseline["avance_programado_pct"], marker="o", label="Avance programado acumulado (%)")
        elif "costo_programado" in baseline.columns:
            plt.plot(baseline["fecha"], baseline["costo_programado"], marker="o", label="Costo programado acumulado")
    else:
        note = note or "No se incluyó línea base programada; la gráfica muestra solo avance real acumulado."

    if len(actual) == 1:
        center = actual["fecha"].iloc[0]
        plt.xlim(center - timedelta(days=3), center + timedelta(days=3))

    plt.title("Curva S de avance semanal")
    plt.xlabel("Fecha")
    plt.ylabel("Avance acumulado / costo")
    plt.grid(True, alpha=0.3)
    plt.legend()
    plt.tight_layout()
    plt.savefig(out_png, dpi=180)
    plt.close()
    return True, note


def add_table(doc: Document, rows: List[List[Any]], widths: Optional[List[float]] = None) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            table.cell(i, j).text = "" if value is None else str(value)
            if i == 0:
                for p in table.cell(i, j).paragraphs:
                    for run in p.runs:
                        run.bold = True
        if widths:
            for j, width in enumerate(widths[: len(row)]):
                table.cell(i, j).width = Inches(width)
    doc.add_paragraph()


def short_date(dt: Optional[datetime]) -> str:
    return dt.strftime("%d/%m/%Y") if dt else "s/f"


def status_week(actual: pd.DataFrame, baseline: pd.DataFrame, reports: List[DailyReport]) -> str:
    serious_incidents = []
    for r in reports:
        for inc in r.incidencias:
            text = " ".join(inc.values()).lower()
            if text and "sin incidencias" not in text and "—" not in text:
                if any(x in text for x in ["alta", "crítico", "critico", "suspensión", "accidente", "retraso"]):
                    serious_incidents.append(text)

    if not actual.empty and not baseline.empty and "avance_programado_pct" in baseline.columns:
        last_actual = actual.dropna(subset=["avance_real_pct"]).tail(1)
        last_prog = baseline.dropna(subset=["avance_programado_pct"]).tail(1)
        if not last_actual.empty and not last_prog.empty:
            diff = float(last_actual["avance_real_pct"].iloc[0]) - float(last_prog["avance_programado_pct"].iloc[0])
            if diff < -5:
                return "Atraso"
            if diff < 0 or serious_incidents:
                return "Con alertas"
            return "En tiempo"
    return "Con información limitada" if serious_incidents else "Sin alertas críticas identificadas"


def collect_key_findings(reports: List[DailyReport], actual: pd.DataFrame, baseline: pd.DataFrame) -> List[str]:
    findings = []
    if reports:
        palabra = "reporte diario" if len(reports) == 1 else "reportes diarios"
        findings.append(f"Se revisaron {len(reports)} {palabra}, del {short_date(reports[0].fecha)} al {short_date(reports[-1].fecha)}.")

    completed = 0
    in_process = 0
    for r in reports:
        for a in r.actividades:
            state = " ".join(a.values()).lower()
            if "complet" in state:
                completed += 1
            elif "proceso" in state or "avance" in state:
                in_process += 1
    if completed or in_process:
        findings.append(f"Se registraron {completed} actividades completadas y {in_process} actividades en proceso durante la semana.")

    last = actual.tail(1) if not actual.empty else pd.DataFrame()
    if not last.empty:
        pct = last["avance_real_pct"].iloc[0] if "avance_real_pct" in last else None
        cost = last["costo_real"].iloc[0] if "costo_real" in last else None
        if pd.notna(pct):
            findings.append(f"El avance acumulado reportado al cierre de la semana es {pct:.1f}%.")
        if pd.notna(cost):
            findings.append(f"El costo ejecutado acumulado al cierre de la semana es ${cost:,.2f}.")

    incident_count = 0
    for r in reports:
        for inc in r.incidencias:
            text = " ".join(inc.values()).lower()
            if text and "sin incidencias" not in text and "—" not in text:
                incident_count += 1
    if incident_count:
        findings.append(f"Se detectaron {incident_count} incidencias relevantes que requieren seguimiento.")
    else:
        findings.append("No se detectaron incidencias relevantes en los reportes revisados.")

    if baseline.empty:
        findings.append("No se proporcionó línea base semanal; por tanto, los atrasos solo pueden identificarse si aparecen explícitos en los RDO.")

    return findings[:6]


def generate_docx(reports: List[DailyReport], baseline: pd.DataFrame, output: Path) -> None:
    reports = sorted(reports, key=lambda r: r.fecha or datetime.min)
    actual = actual_progress_series(reports)

    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    obra = next((r.obra for r in reports if r.obra and r.obra != "No identificado"), "Obra no identificada")
    start = short_date(reports[0].fecha) if reports else "s/f"
    end = short_date(reports[-1].fecha) if reports else "s/f"

    title = doc.add_heading("Reporte semanal de obra", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(obra).bold = True
    doc.add_paragraph(f"Periodo analizado: {start} – {end}")
    doc.add_paragraph(f"Fecha de generación: {datetime.now().strftime('%d/%m/%Y')}")
    doc.add_paragraph()

    estado = status_week(actual, baseline, reports)
    doc.add_heading("1. Resumen ejecutivo", level=1)
    doc.add_paragraph(f"Estado general de la semana: {estado}")
    for item in collect_key_findings(reports, actual, baseline):
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("2. Indicadores clave", level=1)
    avg_personal = None
    vals = [r.personal_total for r in reports if r.personal_total is not None]
    if vals:
        avg_personal = sum(vals) / len(vals)
    avg_hours = None
    hvals = [r.horas_efectivas for r in reports if r.horas_efectivas is not None]
    if hvals:
        avg_hours = sum(hvals) / len(hvals)

    last_actual = actual.tail(1) if not actual.empty else pd.DataFrame()
    avance = "No disponible"
    costo = "No disponible"
    if not last_actual.empty:
        pct = last_actual["avance_real_pct"].iloc[0]
        cst = last_actual["costo_real"].iloc[0]
        if pd.notna(pct):
            avance = f"{pct:.1f}%"
        if pd.notna(cst):
            costo = f"${cst:,.2f}"

    pending_count = sum(len(r.pendientes) for r in reports)
    add_table(
        doc,
        [
            ["Indicador", "Valor"],
            ["Reportes revisados", len(reports)],
            ["Avance acumulado reportado", avance],
            ["Costo ejecutado acumulado", costo],
            ["Personal promedio diario", f"{avg_personal:.1f}" if avg_personal is not None else "No disponible"],
            ["Horas efectivas promedio", f"{avg_hours:.1f}" if avg_hours is not None else "No disponible"],
            ["Pendientes registrados", pending_count],
        ],
    )

    doc.add_heading("3. Avance semanal", level=1)
    rows = [["Fecha", "Reporte", "Actividad", "Cantidad", "Acumulado", "Estado"]]
    for r in reports:
        for a in r.actividades:
            rows.append(
                [
                    short_date(r.fecha),
                    r.no_reporte,
                    a.get("Descripción / Frente", a.get("Descripción", "")),
                    a.get("Cant.", a.get("Cantidad", "")),
                    a.get("Acumulado", ""),
                    a.get("Estado", ""),
                ]
            )
    if len(rows) > 1:
        add_table(doc, rows)
    else:
        doc.add_paragraph("No se encontraron actividades ejecutadas en los reportes revisados.")

    doc.add_heading("4. Curva S", level=1)
    with tempfile.TemporaryDirectory() as tmpdir:
        chart_path = Path(tmpdir) / "curva_s.png"
        ok, note = make_s_curve(actual, baseline, chart_path)
        if ok:
            doc.add_picture(str(chart_path), width=Inches(6.5))
            if note:
                doc.add_paragraph(note)
        else:
            doc.add_paragraph(note)

    doc.add_heading("5. Incidencias y riesgos", level=1)
    inc_rows = [["Fecha", "Tipo", "Urgencia", "Descripción", "Acción tomada"]]
    for r in reports:
        for inc in r.incidencias:
            joined = " ".join(inc.values()).strip()
            if not joined or "sin incidencias" in joined.lower() or joined == "— — — —":
                continue
            inc_rows.append(
                [
                    short_date(r.fecha),
                    inc.get("Tipo", ""),
                    inc.get("Urgencia", ""),
                    inc.get("Descripción", ""),
                    inc.get("Acción tomada", ""),
                ]
            )
    if len(inc_rows) > 1:
        add_table(doc, inc_rows)
    else:
        doc.add_paragraph("No se reportaron incidencias relevantes en la semana.")

    doc.add_heading("6. Materiales, personal y equipo", level=1)
    mat_rows = [["Fecha", "Material", "Proveedor", "Cantidad", "Unidad", "Estado"]]
    for r in reports:
        for m in r.materiales:
            mat_rows.append(
                [
                    short_date(r.fecha),
                    m.get("Material", ""),
                    m.get("Proveedor", ""),
                    m.get("Cant.", m.get("Cantidad", "")),
                    m.get("Unidad", ""),
                    m.get("Estado", ""),
                ]
            )
    if len(mat_rows) > 1:
        add_table(doc, mat_rows)
    else:
        doc.add_paragraph("No se encontraron registros de materiales.")

    doc.add_heading("7. Control presupuestal", level=1)
    if reports and reports[-1].presupuesto:
        keys = list(reports[-1].presupuesto[0].keys())
        rows = [keys]
        for row in reports[-1].presupuesto:
            rows.append([row.get(k, "") for k in keys])
        add_table(doc, rows)
    else:
        doc.add_paragraph("No se encontró tabla de control presupuestal.")

    doc.add_heading("8. Pendientes y acciones recomendadas", level=1)
    pend_rows = [["Fecha RDO", "Tarea / Solicitud", "Responsable", "Fecha límite", "Prioridad sugerida"]]
    for r in reports:
        for p in r.pendientes:
            text = " ".join(p.values()).lower()
            priority = "Alta" if any(w in text for w in ["solicitar", "confirmar", "faltante", "bomba", "entrega"]) else "Media"
            pend_rows.append(
                [
                    short_date(r.fecha),
                    p.get("Tarea / Solicitud", p.get("Tarea", "")),
                    p.get("Responsable", ""),
                    p.get("Fecha límite", p.get("Fecha", "")),
                    priority,
                ]
            )
    if len(pend_rows) > 1:
        add_table(doc, pend_rows)
    else:
        doc.add_paragraph("No se encontraron pendientes registrados.")

    doc.add_heading("9. Limitaciones de la información", level=1)
    if len(reports) < 5:
        palabra = "reporte diario" if len(reports) == 1 else "reportes diarios"
        doc.add_paragraph(f"Solo se recibieron {len(reports)} {palabra}. Para un análisis semanal completo se recomiendan 5 reportes.", style="List Bullet")
    if baseline.empty:
        doc.add_paragraph("No se proporcionó línea base de avance programado. La curva S no permite concluir atraso contra programa.", style="List Bullet")
    doc.add_paragraph("Las conclusiones dependen de la consistencia de los campos capturados en los RDO.", style="List Bullet")

    doc.add_heading("10. Conclusión semanal", level=1)
    if baseline.empty:
        doc.add_paragraph(
            "Con la información disponible, el reporte permite consolidar avance, recursos, incidencias y presupuesto ejecutado. "
            "Para evaluar atrasos y sobrecostos con mayor rigor se requiere incorporar el programa de obra y la línea base financiera."
        )
    else:
        doc.add_paragraph(
            "El desempeño semanal debe evaluarse comparando avance real contra línea base. Las desviaciones detectadas deben convertirse en acciones de recuperación con responsable y fecha límite."
        )

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))


def main() -> None:
    parser = argparse.ArgumentParser(description="Genera reporte semanal de obra desde RDO .docx")
    parser.add_argument("--input", required=True, help="Carpeta con reportes diarios .docx")
    parser.add_argument("--baseline", help="CSV opcional con línea base")
    parser.add_argument("--output", required=True, help="Ruta del .docx de salida")
    args = parser.parse_args()

    input_dir = Path(args.input)
    files = sorted(input_dir.glob("*.docx"))
    if not files:
        raise SystemExit(f"No se encontraron archivos .docx en {input_dir}")

    reports = [parse_report(path) for path in files]
    reports = sorted(reports, key=lambda r: r.fecha or datetime.min)
    baseline = read_baseline(Path(args.baseline)) if args.baseline else pd.DataFrame()
    generate_docx(reports, baseline, Path(args.output))
    print(f"Reporte semanal generado: {args.output}")


if __name__ == "__main__":
    main()
